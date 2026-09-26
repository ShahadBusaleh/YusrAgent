"""Records downloads (Phase 4): Excel list, PDF report, final settlement PDF.

Pure renderers — callers (views.py) pass already-localized labels and
values, so nothing here depends on Streamlit or the interface language
except the `rtl` flag. Arabic is shaped and laid out right-to-left by
fpdf2's HarfBuzz text shaping with the same embedded font the payroll PDF
uses (payroll_report._register_fonts).
"""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Callable

from fpdf import FPDF
from fpdf.fonts import FontFace
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.agents.payroll_report import _register_fonts

FONT = "Payroll"  # family name registered by payroll_report._register_fonts
LOGO = Path(__file__).resolve().parent / "assets" / "yusor_logo_full_brown.png"

BRAND = (70, 46, 36)       # #462E24, the logo brown
BRAND_HEX = "462E24"
TAN = (201, 167, 143)      # #C9A78F
ZEBRA = (247, 243, 238)    # #F7F3EE
LINE = (231, 223, 213)     # #E7DFD5
INK = (36, 26, 20)
MUTED = (78, 70, 61)

# A run of Latin text ("Mona Saleh (EMP-0001)", "(LAW077)"). Parentheses
# only as balanced pairs, so "(المادة 88 (LAW080))" keeps its Arabic ")".
_LATIN_RUN = re.compile(
    r"\([\x20-\x27\x2A-\x7E]*[A-Za-z][\x20-\x27\x2A-\x7E]*\)"
    r"|[A-Za-z](?:[\x20-\x27\x2A-\x7E—–]|\([\x20-\x27\x2A-\x7E]*\))*"
)


def _isolate(text: str) -> str:
    """In RTL text, wrap each Latin run in FSI…PDI so the bidi algorithm
    keeps it in one piece when the cell wraps onto several lines."""

    def wrap(match: re.Match) -> str:
        run = match.group(0)
        core = run.rstrip(" .,:;·—–-")
        return f"\u2068{core}\u2069{run[len(core):]}"

    return _LATIN_RUN.sub(wrap, text)


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def records_xlsx(headers: list[str], rows: list[list], *, rtl: bool, sheet_title: str) -> bytes:
    """One sheet: brand-colored frozen header, autofilter, fitted widths,
    right-to-left when the UI is Arabic."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    ws.sheet_view.rightToLeft = rtl

    ws.append(headers)
    for row in rows:
        ws.append(["" if v is None else v for v in row])

    side = Side(style="thin", color="E7DFD5")
    align = "right" if rtl else "left"
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=BRAND_HEX)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
        cell.border = Border(bottom=side)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(horizontal=align, vertical="top")
            cell.border = Border(bottom=side)
    ws.row_dimensions[1].height = 22

    for index, header in enumerate(headers, start=1):
        values = [str(header)] + [str(r[index - 1] or "") for r in rows]
        width = max(len(v) for v in values)
        ws.column_dimensions[get_column_letter(index)].width = min(max(width + 3, 10), 60)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# PDF helpers
# ---------------------------------------------------------------------------

class _BrandedPDF(FPDF):
    def __init__(self, *, orientation: str, rtl: bool, page_label: str, total_pages: int | None) -> None:
        super().__init__(orientation=orientation, unit="mm", format="A4")
        self.rtl = rtl
        self.page_label = page_label  # "Page {page} of {pages}"
        self.total_pages = total_pages  # known on the second pass (see _two_pass)
        _register_fonts(self)
        self.set_margins(12, 12, 12)
        self.set_auto_page_break(auto=True, margin=16)

    def bidi(self, text: str) -> str:
        return _isolate(text) if self.rtl else text

    @property
    def align(self) -> str:
        return "R" if self.rtl else "L"

    def footer(self) -> None:
        self.set_y(-12)
        self.set_draw_color(*LINE)
        self.line(self.l_margin, self.get_y(), self.w - self.r_margin, self.get_y())
        self.set_font(FONT, "", 8)
        self.set_text_color(*MUTED)
        label = self.page_label.format(page=self.page_no(), pages=self.total_pages or "")
        self.cell(0, 7, label, align="C")

    def logo_header(self, title: str, subtitle: str = "") -> None:
        """Logo on the leading side, title block on the other, brand rule."""
        top = self.get_y()
        logo_w = 26
        logo_x = self.w - self.r_margin - logo_w if self.rtl else self.l_margin
        if LOGO.is_file():
            self.image(str(LOGO), x=logo_x, y=top, w=logo_w)
        text_w = self.w - self.l_margin - self.r_margin - logo_w - 6
        text_x = self.l_margin if self.rtl else self.l_margin + logo_w + 6
        self.set_xy(text_x, top + 2)
        self.set_font(FONT, "B", 16)
        self.set_text_color(*BRAND)
        self.cell(text_w, 9, self.bidi(title), align="L" if self.rtl else "R", new_x="LEFT", new_y="NEXT")
        if subtitle:
            self.set_font(FONT, "", 9)
            self.set_text_color(*MUTED)
            self.cell(text_w, 6, self.bidi(subtitle), align="L" if self.rtl else "R", new_x="LEFT", new_y="NEXT")
        self.set_y(top + logo_w * 518 / 903 + 3)
        self.set_draw_color(*BRAND)
        self.set_line_width(0.6)
        self.line(self.l_margin, self.get_y(), self.w - self.r_margin, self.get_y())
        self.set_line_width(0.2)
        self.ln(4)

    def text_line(self, text: str, *, size: float = 9, bold: bool = False, color=INK, h: float = 5.5) -> None:
        self.set_font(FONT, "B" if bold else "", size)
        self.set_text_color(*color)
        self.set_x(self.l_margin)
        self.multi_cell(0, h, self.bidi(text), align=self.align, new_x="LMARGIN", new_y="NEXT")

    def section_title(self, text: str) -> None:
        self.ln(2)
        self.set_font(FONT, "B", 11)
        self.set_text_color(*BRAND)
        self.set_x(self.l_margin)
        self.cell(0, 7, self.bidi(text), align=self.align, new_x="LMARGIN", new_y="NEXT")
        self.set_draw_color(*TAN)
        self.line(self.l_margin, self.get_y(), self.w - self.r_margin, self.get_y())
        self.ln(1.5)

    def grid(
        self,
        headers: list[str] | None,
        rows: list[list[str]],
        widths: list[float],
        *,
        size: float = 8.5,
        bold_last: bool = False,
        label_cols: tuple[int, ...] = (),
        number_cols: tuple[int, ...] = (),
        padding: float = 1.4,
    ) -> None:
        """Bordered table; column order (and widths) mirrored for RTL.
        Numbers sit on the trailing edge of their column."""
        order = list(reversed(range(len(widths)))) if self.rtl else list(range(len(widths)))
        self.set_font(FONT, "", size)
        self.set_text_color(*INK)
        self.set_draw_color(*LINE)
        head_style = FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=BRAND)
        with self.table(
            col_widths=[widths[i] for i in order],
            text_align="RIGHT" if self.rtl else "LEFT",
            headings_style=head_style,
            first_row_as_headings=headers is not None,
            cell_fill_color=ZEBRA,
            cell_fill_mode="ROWS",
            line_height=size * 0.55,
            padding=padding,
            borders_layout="HORIZONTAL_LINES",
        ) as table:
            if headers is not None:
                head = table.row()
                for i in order:
                    head.cell(self.bidi(headers[i]))
            for n, values in enumerate(rows):
                last = bold_last and n == len(rows) - 1
                row = table.row()
                for i in order:
                    style = None
                    if last:
                        style = FontFace(emphasis="BOLD", fill_color=TAN)
                    elif i in label_cols:
                        style = FontFace(emphasis="BOLD", color=BRAND)
                    text = str(values[i] if values[i] not in (None, "") else "—")
                    align = ("LEFT" if self.rtl else "RIGHT") if i in number_cols else None
                    row.cell(self.bidi(text), style=style, align=align)

    def pairs(self, pairs: list[tuple[str, object]], *, per_row: int = 2) -> None:
        """Label/value pairs, `per_row` pairs side by side."""
        usable = self.w - self.l_margin - self.r_margin
        widths = [usable / per_row * f for _ in range(per_row) for f in (0.36, 0.64)]
        rows = []
        for start in range(0, len(pairs), per_row):
            chunk = list(pairs[start : start + per_row])
            chunk += [(" ", " ")] * (per_row - len(chunk))
            row = []
            for label, value in chunk:
                row += [label, value if value not in (None, "") else ("—" if label else " ")]
            rows.append(row)
        self.grid(None, rows, widths, size=8.5, label_cols=tuple(range(0, 2 * per_row, 2)), padding=1.1)


def _two_pass(render: Callable[[int | None], _BrandedPDF]) -> bytes:
    """Render once to count pages, then again with "Page n of N" filled in
    (fpdf's {nb} alias does not survive HarfBuzz shaping of Arabic)."""
    total = render(None).pages_count
    return bytes(render(total).output())


# ---------------------------------------------------------------------------
# Records report (landscape)
# ---------------------------------------------------------------------------

def records_pdf(
    *,
    title: str,
    subtitle: str,
    filters_title: str,
    filters: list[tuple[str, str]],
    generated: str,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    page_label: str,
    rtl: bool,
) -> bytes:
    def render(total_pages: int | None) -> _BrandedPDF:
        pdf = _BrandedPDF(orientation="L", rtl=rtl, page_label=page_label, total_pages=total_pages)
        pdf.add_page()
        pdf.logo_header(title, subtitle)
        pdf.text_line(filters_title, size=10, bold=True, color=BRAND)
        for label, value in filters:
            pdf.text_line(f"{label}: {value}" if label else value, size=9)
        pdf.text_line(generated, size=8.5, color=MUTED)
        pdf.ln(3)
        usable = pdf.w - pdf.l_margin - pdf.r_margin
        scale = usable / sum(widths)
        pdf.grid(headers, rows, [w * scale for w in widths], size=8)
        return pdf

    return _two_pass(render)


# ---------------------------------------------------------------------------
# Final settlement (portrait)
# ---------------------------------------------------------------------------

def settlement_pdf(doc: dict, *, rtl: bool) -> bytes:
    """`doc` (built by views._settlement_document, all text localized):
    company, company_sub, title, ref, sections [{title, pairs}], calc
    {title, headers, rows}, notes [str], approval {title, pairs},
    signatures {ack, name, signature, date, blocks [title, title]},
    page_label."""
    return _two_pass(lambda total: _settlement(doc, rtl=rtl, total_pages=total))


def _settlement(doc: dict, *, rtl: bool, total_pages: int | None) -> _BrandedPDF:
    pdf = _BrandedPDF(orientation="P", rtl=rtl, page_label=doc["page_label"], total_pages=total_pages)
    pdf.add_page()
    pdf.logo_header(doc["company"], doc["company_sub"])

    pdf.set_font(FONT, "B", 17)
    pdf.set_text_color(*INK)
    pdf.cell(0, 10, pdf.bidi(doc["title"]), align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font(FONT, "", 9)
    pdf.set_text_color(*MUTED)
    pdf.cell(0, 5, pdf.bidi(doc["ref"]), align="C", new_x="LMARGIN", new_y="NEXT")

    usable = pdf.w - pdf.l_margin - pdf.r_margin
    for section in doc["sections"]:
        pdf.section_title(section["title"])
        pdf.pairs(section["pairs"])

    calc = doc["calc"]
    pdf.section_title(calc["title"])
    pdf.grid(
        calc["headers"],
        calc["rows"],
        [usable * w for w in (0.24, 0.40, 0.18, 0.18)],
        size=8.5,
        bold_last=True,
        number_cols=(3,),
        padding=1.1,
    )
    pdf.ln(1)
    for note in doc["notes"]:
        pdf.text_line(note, size=8, color=MUTED, h=4.5)

    approval = doc["approval"]
    pdf.section_title(approval["title"])
    pdf.pairs(approval["pairs"])

    sig = doc["signatures"]
    if pdf.get_y() + 42 > pdf.h - pdf.b_margin:  # acknowledgement + both signature blocks
        pdf.add_page()
    pdf.ln(3)
    pdf.text_line(sig["ack"], size=9)
    pdf.ln(1)
    block_w = (usable - 10) / 2
    top = pdf.get_y()
    blocks = list(reversed(sig["blocks"])) if rtl else sig["blocks"]
    for n, block_title in enumerate(blocks):
        x = pdf.l_margin + n * (block_w + 10)
        pdf.set_xy(x, top)
        pdf.set_font(FONT, "B", 10)
        pdf.set_text_color(*BRAND)
        pdf.cell(block_w, 7, pdf.bidi(block_title), align=pdf.align, new_x="LEFT", new_y="NEXT")
        pdf.set_font(FONT, "", 9)
        pdf.set_text_color(*INK)
        for label in (sig["name"], sig["signature"], sig["date"]):
            y = pdf.get_y() + 8
            pdf.set_xy(x, y - 5)
            pdf.cell(block_w, 5, pdf.bidi(label), align=pdf.align)
            pdf.set_draw_color(*MUTED)
            if rtl:
                pdf.line(x, y, x + block_w - 22, y)
            else:
                pdf.line(x + 22, y, x + block_w, y)
            pdf.set_y(y)
    return pdf
